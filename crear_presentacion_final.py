"""Genera una presentación nueva y editable de los resultados del Proyecto 2."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "entregables/final/Proyecto2_Resultados_Jorge_Lopez.pptx"
BG = RGBColor(247, 248, 243)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(30, 52, 41)
GREEN = RGBColor(47, 105, 70)
MUTED = RGBColor(92, 109, 97)
LINE = RGBColor(215, 225, 214)
BLUE = RGBColor(82, 125, 145)
ORANGE = RGBColor(217, 139, 78)
VIOLET = RGBColor(142, 116, 165)
PALE = RGBColor(233, 241, 232)
W, H = 13.333, 7.5


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def oval(slide, x, y, w, h, color, transparency=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency
    shape.line.fill.background()
    return shape


def txt(slide, value, x, y, w, h, size=18, color=INK, bold=False,
        font="Aptos", align=None, valign=MSO_ANCHOR.TOP, margin=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.text = line
        para.space_after = Pt(0)
        para.line_spacing = 1.12
        if align is not None:
            para.alignment = align
        for run in para.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return box


def picture(slide, filename, x, y, w, h):
    path = ROOT / filename
    with Image.open(path) as image:
        width, height = image.size
    scale = min(w / width, h / height)
    pw, ph = width * scale, height * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w-pw)/2),
                                    Inches(y + (h-ph)/2), width=Inches(pw), height=Inches(ph))


def base(number, kicker, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    txt(slide, "CC3084  /  PROYECTO 2  /  TEMA 22", .55, .22, 7, .24, 9, GREEN, True)
    rect(slide, .55, .57, 12.23, .018, LINE)
    txt(slide, kicker.upper(), .66, .75, 10, .25, 10, GREEN, True)
    txt(slide, title, .66, 1.03, 12, .62, 30, INK, True, "Georgia")
    rect(slide, .55, 7.1, 12.23, .014, LINE)
    txt(slide, "Jorge López · 221038", .66, 7.19, 6, .18, 9, MUTED)
    txt(slide, f"{number:02d} / 10", 11.5, 7.19, 1.2, .18, 9, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, label, number, detail, accent=GREEN):
    rect(slide, x, y, w, h, WHITE, LINE, True)
    rect(slide, x, y, .065, h, accent)
    txt(slide, label.upper(), x+.25, y+.22, w-.5, .25, 10, MUTED, True)
    txt(slide, number, x+.25, y+.58, w-.5, .65, 31, accent, False, "Georgia")
    txt(slide, detail, x+.25, y+1.25, w-.5, h-1.4, 12, INK)


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
prs.core_properties.author = "Jorge López"
prs.core_properties.title = "Proyecto 2 · Resultados · Clasificación de coágulos"
prs.core_properties.subject = "CC3084 Data Science · Tema 22 · Mayo Clinic STRIP AI"

# 1. Portada.
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
oval(s, 9.1, .45, 3.35, 3.35, RGBColor(194, 154, 146), 35)
oval(s, 10.8, 2.5, 1.75, 1.75, RGBColor(151, 110, 127), 25)
txt(s, "UNIVERSIDAD DEL VALLE DE GUATEMALA  /  CC3084", .7, .5, 10, .3, 11, GREEN, True)
rect(s, .7, 2.0, .9, .05, GREEN)
txt(s, "Clasificación del origen\nde coágulos en ACV", .7, 2.25, 10.4, 1.55, 36, INK, True, "Georgia")
txt(s, "Resultados del Proyecto 2 · Tema 22", .72, 4.1, 9, .4, 19, GREEN)
txt(s, "Jorge López  ·  221038", .72, 5.65, 6, .35, 17, INK, True)
txt(s, "12 de octubre de 2026", .72, 6.1, 6, .28, 12, MUTED)
rect(s, .7, 7.05, 11.9, .025, LINE)

# 2. Ampliación crucial de datos.
s = base(2, "Del piloto al conjunto completo", "De 6 imágenes a 754 verificadas")
card(s, .68, 1.86, 4.0, 2.35, "Primera entrega", "6", "archivos visuales derivados; solo 3 tenían etiqueta vinculable al CSV", BLUE)
txt(s, "→", 5.0, 2.45, .9, .7, 35, GREEN, True, align=PP_ALIGN.CENTER)
card(s, 6.05, 1.86, 6.58, 2.35, "Segunda fase", "754 / 754", "imágenes etiquetadas de entrenamiento en PNG reducido · 632 pacientes · 11 centros", GREEN)
rect(s, .68, 4.53, 11.95, 1.55, PALE, None, True)
txt(s, "Qué cambió en el análisis", .97, 4.8, 6, .3, 17, GREEN, True)
txt(s, "Rehicimos el EDA visual y la máscara de tejido sobre todas las imágenes disponibles para el entrenamiento. El EDA anterior era un piloto y no podía representar al conjunto completo.", .97, 5.2, 11.1, .7, 16, INK)
txt(s, "Alcance exacto: copia PNG reducida de las láminas etiquetadas; no son todos los TIFF originales ni el conjunto oculto de prueba.", .73, 6.35, 11.8, .42, 11, MUTED)

# 3. Pregunta y objetivo.
s = base(3, "Problema y objetivo", "¿Puede la imagen sugerir el origen?")
rect(s, .69, 1.84, 11.9, 1.18, PALE, None, True)
txt(s, "Pregunta de investigación", 1.0, 2.02, 3.1, .3, 12, GREEN, True)
txt(s, "¿Generalizan a pacientes nuevos los modelos basados en color y textura de regiones con tejido?", 1.0, 2.36, 10.9, .59, 17, INK, True)
rect(s, .69, 3.31, 5.78, 2.5, WHITE, LINE, True)
rect(s, 6.78, 3.31, 5.78, 2.5, WHITE, LINE, True)
oval(s, 1.04, 3.67, .72, .72, BLUE)
oval(s, 7.13, 3.67, .72, .72, ORANGE)
txt(s, "CE", 1.21, 3.87, .4, .2, 13, WHITE, True, align=PP_ALIGN.CENTER)
txt(s, "LAA", 7.2, 3.87, .59, .2, 11, WHITE, True, align=PP_ALIGN.CENTER)
txt(s, "Cardioembólico", 1.0, 4.57, 4.8, .35, 21, INK, True, "Georgia")
txt(s, "Origen asociado con una fuente cardíaca.", 1.0, 5.1, 4.9, .5, 15, MUTED)
txt(s, "Arteria grande", 7.1, 4.57, 4.8, .35, 21, INK, True, "Georgia")
txt(s, "Ateroesclerosis de grandes arterias.", 7.1, 5.1, 4.9, .5, 15, MUTED)
txt(s, "Meta: comparar tres algoritmos por paciente y comunicar sus límites en una aplicación interactiva.", .74, 6.22, 11.7, .43, 16, GREEN, True)

# 4. EDA general.
s = base(4, "EDA visual completo", "El conjunto es desigual y variable")
rect(s, .68, 1.86, 4.14, 4.48, WHITE, LINE, True)
txt(s, "PACIENTES POR CLASE", .98, 2.15, 3.5, .25, 11, MUTED, True)
txt(s, "CE", .98, 2.68, .5, .27, 15, INK, True)
rect(s, 1.6, 2.69, 2.62, .22, PALE, None, True)
rect(s, 1.6, 2.69, 1.90, .22, BLUE, None, True)
txt(s, "457", 4.17, 2.68, .5, .27, 15, INK, True, align=PP_ALIGN.RIGHT)
txt(s, "LAA", .98, 3.28, .6, .27, 15, INK, True)
rect(s, 1.6, 3.29, 2.62, .22, PALE, None, True)
rect(s, 1.6, 3.29, .72, .22, ORANGE, None, True)
txt(s, "175", 4.17, 3.28, .5, .27, 15, INK, True, align=PP_ALIGN.RIGHT)
txt(s, "12.45 %", .98, 4.28, 3.4, .6, 29, GREEN, False, "Georgia")
txt(s, "mediana de área candidata a tejido", .98, 4.98, 3.3, .55, 13, MUTED)
txt(s, "7 imágenes bajo 1 %: revisadas una por una.", .98, 5.78, 3.4, .42, 12, INK)
rect(s, 5.09, 1.86, 7.5, 4.48, WHITE, LINE, True)
picture(s, "resultados/figuras/reducidas_cobertura_tejido.png", 5.27, 2.1, 7.15, 3.52)
txt(s, "La cobertura varía entre centros; no prueba diferencias biológicas CE/LAA.", 5.35, 5.75, 7.1, .45, 12, MUTED)
txt(s, "Fuente: EDA de las 754 imágenes reducidas; medidas por imagen, no por paciente.", .73, 6.59, 11.7, .26, 10, MUTED)

# 5. Tejido y regiones.
s = base(5, "Preprocesamiento", "Separar tejido y fondo")
rect(s, .68, 1.83, 7.28, 3.77, WHITE, LINE, True)
picture(s, "resultados/figuras/reducidas_comparacion_mascaras.png", .85, 2.04, 6.94, 3.08)
txt(s, "La regla inicial incluía fondos teñidos; la máscara nueva usa color y textura.", .89, 5.14, 6.86, .33, 11, MUTED)
card(s, 8.24, 1.83, 4.32, 1.74, "Regiones", "5,933", "recortes de 256 × 256 px", GREEN)
rect(s, 8.24, 3.79, 4.32, 1.81, PALE, None, True)
txt(s, "750 de 754 imágenes", 8.48, 4.04, 3.9, .41, 20, GREEN, True, "Georgia")
txt(s, "630 de 632 pacientes con regiones; los 2 restantes estaban en entrenamiento.", 8.48, 4.55, 3.78, .73, 13, INK)
rect(s, .69, 5.88, 11.87, .74, PALE, None, True)
txt(s, "Control de calidad: 61 pares imagen/máscara revisados; las 7 imágenes bajo 1 % se inspeccionaron visualmente.", .98, 6.08, 11.2, .37, 13, INK)

# 6. Diseño de la evaluación.
s = base(6, "Diseño experimental", "Partición por paciente, sin fuga")
for x, accent, name, num, detail in [
    (.7, BLUE, "ENTRENAMIENTO", "440", "438 con regiones"),
    (4.93, GREEN, "VALIDACIÓN", "97", "elige el modelo"),
    (9.16, ORANGE, "PRUEBA", "95", "evaluación final")]:
    card(s, x, 1.86, 3.45, 2.37, name, num, detail, accent)
rect(s, .7, 4.57, 11.9, 1.34, WHITE, LINE, True)
txt(s, "57 medidas por región", 1.03, 4.86, 3.13, .35, 19, INK, True, "Georgia")
txt(s, "→", 4.32, 4.82, .5, .42, 25, GREEN, True)
txt(s, "114 variables por paciente", 4.91, 4.86, 4.75, .35, 19, GREEN, True, "Georgia")
txt(s, "Color, histogramas, entropía y textura; centro e ID no entran al modelo.", 1.03, 5.38, 10.9, .31, 13, MUTED)
txt(s, "Regresión logística  ·  Bosque aleatorio  ·  SVM RBF", .85, 6.25, 11.7, .37, 17, INK, True, align=PP_ALIGN.CENTER)
txt(s, "Búsqueda de hiperparámetros: 4 pliegues solo en entrenamiento · selección por AUC en validación.", .85, 6.69, 11.7, .23, 10, MUTED, align=PP_ALIGN.CENTER)

# 7. Comparación de modelos.
s = base(7, "Resultados", "El ganador no generalizó bien")
rect(s, .7, 1.85, 11.9, 3.41, WHITE, LINE, True)
txt(s, "AUC ROC", 1.02, 2.12, 3, .34, 17, INK, True)
models = [("Regresión logística", .608, .528, BLUE), ("Bosque aleatorio", .641, .522, ORANGE),
          ("SVM RBF", .575, .461, VIOLET), ("Siempre CE", .500, .500, MUTED)]
for i,(name,val,test,color) in enumerate(models):
    y=2.65+i*.58
    txt(s, name, 1.02, y, 2.72, .25, 14, INK, i==1)
    rect(s, 3.82, y+.02, 5.8, .20, PALE, None, True)
    rect(s, 3.82, y+.02, 5.8*test, .20, color, None, True)
    txt(s, f"Prueba {test:.3f}", 9.79, y-.01, 2.22, .28, 13, color, True)
    txt(s, f"Val. {val:.3f}", 11.66, y-.01, .79, .28, 11, MUTED, align=PP_ALIGN.RIGHT)
rect(s, .7, 5.61, 11.9, .90, PALE, None, True)
txt(s, "Bosque elegido por validación: AUC 0.641; en prueba: 0.522", .97, 5.87, 11.25, .39, 20, GREEN, True, "Georgia")
txt(s, "AUC 0.5 ≈ azar. No se cambia el ganador después de mirar la prueba.", .75, 6.70, 11.7, .28, 12, MUTED)

# 8. Errores.
s = base(8, "Errores del bosque aleatorio", "Solo identificó 5 de 26 pacientes LAA")
txt(s, "Matriz de confusión en 95 pacientes de prueba", .74, 1.79, 9, .34, 17, INK, True)
txt(s, "Predicción CE", 4.03, 2.40, 2.4, .3, 13, MUTED, True, align=PP_ALIGN.CENTER)
txt(s, "Predicción LAA", 7.03, 2.40, 2.4, .3, 13, MUTED, True, align=PP_ALIGN.CENTER)
txt(s, "Real CE", 1.0, 3.30, 2.1, .4, 17, BLUE, True)
txt(s, "Real LAA", 1.0, 4.66, 2.1, .4, 17, ORANGE, True)
for x,y,n,label,fill in [(3.81,2.75,"54","aciertos CE",PALE), (6.8,2.75,"15","falsos LAA",RGBColor(252,235,222)),
                         (3.81,4.10,"21","LAA omitidos",RGBColor(252,235,222)), (6.8,4.10,"5","aciertos LAA",PALE)]:
    rect(s, x, y, 2.62, 1.08, fill, None, True)
    txt(s, n, x+.18, y+.13, 1.0, .49, 28, GREEN if fill==PALE else ORANGE, True, "Georgia")
    txt(s, label, x+.98, y+.43, 1.48, .31, 12, INK)
rect(s, 10.1, 2.75, 2.34, 2.43, WHITE, LINE, True)
txt(s, "19.2 %", 10.34, 3.13, 1.83, .55, 25, ORANGE, True, "Georgia")
txt(s, "sensibilidad LAA", 10.34, 3.79, 1.79, .62, 13, MUTED)
txt(s, "La exactitud total (62.1 %) fue inferior a la línea base que siempre predice CE (72.6 %).", .83, 5.73, 11.55, .43, 17, INK, True)
txt(s, "Interpretación: estas características simples de PNG reducidos no sustentan una clasificación clínica.", .83, 6.44, 11.55, .47, 14, GREEN)

# 9. Aplicación.
s = base(9, "Aplicación interactiva", "Explorar, comparar y probar imágenes")
rect(s, .69, 1.83, 5.86, 4.40, WHITE, LINE, True)
rect(s, 6.78, 1.83, 5.86, 4.40, WHITE, LINE, True)
picture(s, "resultados/captura_app_caracteristicas.png", .87, 2.01, 5.5, 3.65)
picture(s, "resultados/captura_app_prediccion.png", 6.96, 2.01, 5.5, 3.65)
txt(s, "114 características filtrables", 1.02, 5.79, 5.18, .31, 16, GREEN, True)
txt(s, "Tres respuestas por paciente", 7.1, 5.79, 5.18, .31, 16, GREEN, True)
rect(s, .72, 6.42, 11.89, .45, PALE, None, True)
txt(s, "Preprocesamiento automático · métricas interactivas · opción de ocultar el rendimiento", .99, 6.54, 11.35, .22, 12, INK, align=PP_ALIGN.CENTER)

# 10. Cierre y bibliografía.
s = base(10, "Conclusiones", "Flujo completo; modelo exploratorio")
for i,(head,body) in enumerate([
    ("Datos", "Conseguimos las 754 imágenes etiquetadas en PNG reducido y rehicimos el EDA; el piloto de seis imágenes quedó delimitado."),
    ("Método", "Particionamos 632 pacientes sin mezclar imágenes, seleccionamos tejido y comparamos tres modelos con prueba reservada."),
    ("Resultado", "El bosque elegido obtuvo AUC 0.522 y detectó 5/26 LAA en prueba. La app demuestra el flujo, no un uso clínico.")]):
    y=1.88+i*1.25
    rect(s, .72, y, 11.85, 1.07, WHITE, LINE, True)
    txt(s, f"0{i+1}", .98, y+.22, .6, .46, 23, GREEN, True, "Georgia")
    txt(s, head, 1.67, y+.17, 2.0, .36, 18, INK, True)
    txt(s, body, 3.64, y+.18, 8.49, .65, 14, MUTED)
rect(s, .72, 5.89, 11.85, .89, PALE, None, True)
txt(s, "Referencias principales", .99, 6.02, 3.0, .26, 11, GREEN, True)
txt(s, "Komura & Ishikawa (2018) · Patel et al. (2023) · Lucero-Garófano et al. (2025)\nDOI y bibliografía APA completos en el informe final.", .99, 6.32, 11.18, .39, 10, INK)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"Guardada: {OUT} ({len(prs.slides)} diapositivas)")
