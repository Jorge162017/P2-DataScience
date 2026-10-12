"""Exporta las formas, textos e imágenes del PPTX del proyecto a PDF."""

from pathlib import Path

import fitz
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "presentacion/Proyecto2_Resultados_Jorge_Lopez.pptx"
PDF = PPTX.with_suffix(".pdf")
FONTS = Path("/System/Library/Fonts/Supplemental")
FONT_FILES = {
    "sans": FONTS / "Arial.ttf",
    "sans-bold": FONTS / "Arial Bold.ttf",
    "serif": FONTS / "Georgia.ttf",
    "serif-bold": FONTS / "Georgia Bold.ttf",
}
EMU_PER_POINT = 12700


def rgb(color, fallback=(0, 0, 0)):
    try:
        raw = color.rgb
        if raw is None:
            return fallback
        return tuple(channel / 255 for channel in raw)
    except (AttributeError, TypeError, ValueError):
        return fallback


def rect_for(shape):
    return fitz.Rect(*(v / EMU_PER_POINT for v in
                       (shape.left, shape.top, shape.left + shape.width,
                        shape.top + shape.height)))


def draw_shape(page, shape):
    box = rect_for(shape)
    try:
        fill = rgb(shape.fill.fore_color) if shape.fill.type else None
    except (AttributeError, TypeError):
        fill = None
    try:
        stroke = rgb(shape.line.color) if shape.line.fill.type else None
    except (AttributeError, TypeError):
        stroke = None
    width = shape.line.width / EMU_PER_POINT if stroke else 0
    name = str(shape.auto_shape_type)
    if "OVAL" in name:
        page.draw_oval(box, color=stroke, fill=fill, width=width)
    else:
        page.draw_rect(box, color=stroke, fill=fill, width=width)


def draw_text(page, shape):
    frame = shape.text_frame
    box = rect_for(shape)
    box.x0 += frame.margin_left / EMU_PER_POINT
    box.x1 -= frame.margin_right / EMU_PER_POINT
    box.y0 += frame.margin_top / EMU_PER_POINT
    box.y1 -= frame.margin_bottom / EMU_PER_POINT
    if not shape.text:
        return
    paragraph = next((p for p in frame.paragraphs if p.runs), None)
    if paragraph is None:
        return
    font = paragraph.runs[0].font
    family = "serif" if font.name == "Georgia" else "sans"
    font_name = family + ("-bold" if font.bold else "")
    size = font.size.pt if font.size else 14
    font_id = font_name.replace("-", "_")
    page.insert_font(fontname=font_id, fontfile=str(FONT_FILES[font_name]))
    align = int(paragraph.alignment) if paragraph.alignment is not None else 0
    # The same textbox may contain multiple PowerPoint paragraphs.
    available = fitz.Rect(box.x0, box.y0, box.x1, box.y1 + 3)
    for smaller in (size, size - .5, size - 1, size - 1.5, size - 2,
                    size - 2.5, size - 3, size - 4):
        if smaller <= 0:
            break
        spare = page.insert_textbox(available, shape.text,
                                    fontsize=smaller, fontname=font_id,
                                    color=rgb(font.color), align=align,
                                    lineheight=1.08)
        if spare >= 0:
            break
    if spare < 0:
        print(f"AVISO: texto no cabía: {shape.text[:70]}")


def main():
    prs = Presentation(PPTX)
    doc = fitz.open()
    for slide in prs.slides:
        page = doc.new_page(width=prs.slide_width / EMU_PER_POINT,
                            height=prs.slide_height / EMU_PER_POINT)
        page.draw_rect(page.rect, color=None, fill=rgb(slide.background.fill.fore_color))
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                page.insert_image(rect_for(shape), stream=shape.image.blob)
            elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                draw_shape(page, shape)
            elif shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                draw_text(page, shape)
    doc.set_metadata({"title": "Proyecto 2 · Resultados · Clasificación de coágulos",
                      "author": "Jorge López"})
    doc.save(PDF, garbage=4, deflate=True)
    print(f"Guardado: {PDF} ({len(doc)} páginas)")


if __name__ == "__main__":
    main()
